"""
Assignment 2 - Question 2: Batch Gradebook Generator & Excel Formatter
Subject: Python Laboratory (DJS23APC203L)
Module Focus: Structured Spreadsheet Automation with openpyxl
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def generate_gradebook(output_filename: str = "AIML_Batch_Grades.xlsx"):
    """
    Creates a formatted Excel gradebook with computed fields,
    styles, and summary class metrics.
    """
    # 1. Programmatically instantiate a new workbook and select active sheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Semester Results"

    # 2. Populate Header Row
    headers = [
        "Roll No",
        "Student Name",
        "Lab Marks (out of 50)",
        "Theory Marks (out of 50)",
        "Total Marks (out of 100)",
        "Result"
    ]
    ws.append(headers)

    # 3. Sample student records (at least 5 rows)
    student_records = [
        (101, "Aarav Sharma", 42, 38),
        (102, "Diya Patel", 48, 45),
        (103, "Rohan Kulkarni", 18, 16),   # Total: 34 -> Fail
        (104, "Ananya Iyer", 35, 41),
        (105, "Kabir Mehta", 22, 17),    # Total: 39 -> Fail
        (106, "Sneha Joshi", 40, 39),
    ]

    # 4. Programmatically calculate Total Marks and Result
    for record in student_records:
        roll_no, name, lab_marks, theory_marks = record
        total_marks = lab_marks + theory_marks
        result = "Pass" if total_marks > 40 else "Fail"

        ws.append([roll_no, name, lab_marks, theory_marks, total_marks, result])

    # 5. Styling & Aggregations
    header_font = Font(name="Calibri", size=11, bold=True, color="000000")
    header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")  # Light gray
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )

    # Apply header styling
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Format data rows & calculate scores for class average
    total_scores = []
    start_row = 2
    end_row = start_row + len(student_records) - 1

    for row_idx in range(start_row, end_row + 1):
        total_val = ws.cell(row=row_idx, column=5).value
        total_scores.append(total_val)

        # Center alignment and borders for data cells
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.border = thin_border
            if col_idx in [1, 3, 4, 5, 6]:
                c.alignment = center_align

    # Calculate class average
    class_average = round(sum(total_scores) / len(total_scores), 2)

    # Append summary row for Class Average
    summary_row_idx = end_row + 1
    ws.cell(row=summary_row_idx, column=1, value="Class Average")
    ws.merge_cells(start_row=summary_row_idx, start_column=1, end_row=summary_row_idx, end_column=4)

    avg_cell = ws.cell(row=summary_row_idx, column=5, value=class_average)
    ws.cell(row=summary_row_idx, column=6, value="")

    # Style summary row
    summary_font = Font(name="Calibri", size=11, bold=True, color="1A365D")
    summary_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=summary_row_idx, column=col_idx)
        c.font = summary_font
        c.fill = summary_fill
        c.border = thin_border
    ws.cell(row=summary_row_idx, column=1).alignment = Alignment(horizontal="right", vertical="center")
    avg_cell.alignment = center_align

    # Adjust column widths automatically for clean display
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Save workbook in the same folder as the script
    script_dir = os.path.dirname(os.path.abspath(__file__))
    save_path = os.path.join(script_dir, output_filename)
    wb.save(save_path)

    print("=" * 60)
    print("         BATCH GRADEBOOK GENERATOR EXECUTED         ")
    print("=" * 60)
    print(f"[+] Total Student Records Processed : {len(student_records)}")
    print(f"[+] Class Average Score             : {class_average} / 100")
    print(f"[+] Workbook successfully saved at  : {save_path}")
    print("=" * 60)


if __name__ == "__main__":
    generate_gradebook("AIML_Batch_Grades.xlsx")